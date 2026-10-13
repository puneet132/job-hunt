#!/usr/bin/env python3
"""Build a one-page, ATS-friendly resume (DOCX, optionally PDF) from a JSON file.

Usage:
    python build_resume.py resume.json --out Jane_Doe_Product_Owner_Resume.docx [--pdf]

The JSON shape is documented in ../assets/resume-example.json.
Requires python-docx (pip install python-docx). PDF export needs LibreOffice (soffice) on PATH.
"""
import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: F401  (kept for compatibility)
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Inches, RGBColor
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")

FONT = "Calibri"
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_base_style(doc, size):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.05


def para(doc, text="", size=None, bold=False, italic=False, color=None, align=None, after=0, before=0):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "444444")
    bdr.append(b)
    pPr.append(bdr)


def section(doc, title):
    p = para(doc, title.upper(), size=10.5, bold=True, before=7, after=3)
    p.runs[0].font.color.rgb = INK
    bottom_border(p)


def left_right(doc, left, right, width_in, left_bold=False, left_italic=False, size=None, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(width_in), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run(left)
    r1.bold = left_bold
    r1.italic = left_italic
    r2 = p.add_run("\t" + (right or ""))
    r2.italic = left_italic
    for r in (r1, r2):
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def bullet(doc, text, size):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(1)
    # Support **bold** fragments inside bullets.
    parts = text.split("**")
    for i, part in enumerate(parts):
        if part:
            r = p.add_run(part)
            r.bold = i % 2 == 1
            r.font.size = Pt(size)
    return p


def build(data, out, body_size=10.0):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    sec.left_margin = sec.right_margin = Inches(0.55)
    sec.top_margin = Inches(0.45)
    sec.bottom_margin = Inches(0.4)
    width = 8.27 - 1.1
    set_base_style(doc, body_size)

    para(doc, data["name"], size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    if data.get("headline"):
        para(doc, data["headline"], size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    contact = " | ".join(x for x in data.get("contact", []) if x)
    para(doc, contact, size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)

    if data.get("summary"):
        section(doc, "Summary")
        para(doc, data["summary"], size=body_size, bold=data.get("summary_bold", False))

    if data.get("experience"):
        section(doc, "Experience")
        for job in data["experience"]:
            left_right(doc, f'{job["title"]} at {job["company"]}', job.get("dates", ""), width,
                       left_bold=True, size=body_size + 1)
            if job.get("context") or job.get("location"):
                left_right(doc, job.get("context", ""), job.get("location", ""), width,
                           left_italic=True, size=body_size - 1, color=MUTED)
            for b in job.get("bullets", []):
                bullet(doc, b, body_size)

    if data.get("projects"):
        section(doc, data.get("projects_title", "Projects"))
        for b in data["projects"]:
            bullet(doc, b, body_size)

    if data.get("skills"):
        section(doc, "Skills")
        for label, items in data["skills"].items():
            p = doc.add_paragraph()
            r = p.add_run(label + ": ")
            r.bold = True
            p.add_run(items)

    if data.get("education") or data.get("awards"):
        section(doc, data.get("education_title", "Education & Awards"))
        for ed in data.get("education", []):
            left_right(doc, ed["degree"], ed.get("dates", ""), width, left_bold=True)
            if ed.get("note"):
                para(doc, ed["note"], italic=True, size=body_size - 1, color=MUTED)
        if data.get("awards"):
            p = doc.add_paragraph()
            r = p.add_run("Awards: ")
            r.bold = True
            p.add_run("; ".join(data["awards"]))

    doc.save(out)
    return out


def to_pdf(docx_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found; skipped PDF export.", file=sys.stderr)
        return None
    outdir = str(Path(docx_path).resolve().parent)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, str(docx_path)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = Path(docx_path).with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json_file")
    ap.add_argument("--out", required=True, help="Output .docx path")
    ap.add_argument("--pdf", action="store_true", help="Also export a PDF (needs LibreOffice)")
    ap.add_argument("--size", type=float, default=10.0, help="Body font size in pt (default 10)")
    a = ap.parse_args()
    data = json.loads(Path(a.json_file).read_text(encoding="utf-8"))
    out = build(data, a.out, a.size)
    print(f"Wrote {out}")
    if a.pdf:
        pdf = to_pdf(out)
        if pdf:
            print(f"Wrote {pdf}")
            try:
                pages = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
                n = [l for l in pages.splitlines() if l.startswith("Pages:")]
                if n and n[0].split()[-1] != "1":
                    print("Warning: resume is longer than one page; shorten bullets or use --size 9.5", file=sys.stderr)
            except FileNotFoundError:
                pass


if __name__ == "__main__":
    main()
